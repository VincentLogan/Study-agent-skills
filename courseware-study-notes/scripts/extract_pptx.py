#!/usr/bin/env python3
"""Extract PPTX source facts into layout-aware JSON for course-note generation."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any, Iterable

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
except ModuleNotFoundError:
    print("Missing python-pptx. Run: python -m pip install -r requirements.txt", file=sys.stderr)
    raise SystemExit(2)

EMU_PER_INCH = 914400
EMU_PER_POINT = 12700


def inches(value: int | None) -> float | None:
    return round(value / EMU_PER_INCH, 4) if value is not None else None


def points(value: Any) -> float | None:
    try:
        return round(int(value) / EMU_PER_POINT, 2) if value is not None else None
    except (TypeError, ValueError):
        return None


def walk(shapes: Iterable[Any]) -> Iterable[Any]:
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)
        else:
            yield shape


def is_title(shape: Any) -> bool:
    try:
        return shape.is_placeholder and shape.placeholder_format.type in {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    except (AttributeError, ValueError):
        return False


def paragraph_data(paragraph: Any) -> dict[str, Any]:
    runs, sizes = [], []
    for run in paragraph.runs:
        size = points(run.font.size)
        if size is not None:
            sizes.append(size)
        runs.append({"text": run.text, "font_size_pt": size, "bold": run.font.bold, "italic": run.font.italic})
    return {"text": paragraph.text, "level": paragraph.level, "font_size_pt": max(sizes) if sizes else None, "runs": runs}


def slide_data(slide: Any, number: int, mode: str) -> dict[str, Any]:
    blocks, non_text = [], 0
    for index, shape in enumerate(walk(slide.shapes), 1):
        base = {"shape_index": index, "shape_name": shape.name, "x_in": inches(shape.left), "y_in": inches(shape.top), "width_in": inches(shape.width), "height_in": inches(shape.height)}
        if getattr(shape, "has_table", False):
            rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
            blocks.append(base | {"kind": "table", "rows": rows, "text": "\n".join("\t".join(row) for row in rows)})
        elif getattr(shape, "has_text_frame", False) and shape.text.strip():
            paragraphs = [paragraph_data(p) for p in shape.text_frame.paragraphs if p.text.strip()]
            blocks.append(base | {"kind": "text", "is_title_placeholder": is_title(shape), "text": "\n".join(p["text"] for p in paragraphs), "paragraphs": paragraphs})
        else:
            non_text += 1
    blocks.sort(key=lambda block: (block["y_in"] or 0, block["x_in"] or 0, block["shape_index"]))
    title = next((block["text"] for block in blocks if block.get("is_title_placeholder")), None)
    if not title:
        candidates = [block for block in blocks if block["kind"] == "text"]
        if candidates:
            title = max(candidates, key=lambda block: max((p["font_size_pt"] or 0 for p in block["paragraphs"]), default=0))["text"]
    try:
        notes = (slide.notes_slide.notes_text_frame.text or "").strip()
    except (AttributeError, ValueError):
        notes = ""
    native_text = "\n\n".join(block["text"] for block in blocks)
    index = {
        "slide_number": number,
        "title_candidate": title,
        "native_text": native_text,
        "native_text_characters": len(native_text),
        "speaker_notes": notes,
        "non_text_shape_count": non_text,
    }
    if mode == "index":
        return index
    return index | {"text_blocks": blocks}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--mode", choices=("index", "full"), default="full", help="Use index for compact course mapping; full retains layout metadata.")
    args = parser.parse_args()
    if args.input.suffix.lower() != ".pptx" or not args.input.is_file():
        parser.error("--input must name an existing .pptx file")
    deck = Presentation(args.input)
    data = {"source_file": str(args.input.resolve()), "mode": args.mode, "slide_count": len(deck.slides), "slides": [slide_data(slide, index, args.mode) for index, slide in enumerate(deck.slides, 1)]}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
